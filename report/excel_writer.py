import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.common_rules import (
    get_heading_check_report_name,
    safe_sheet_name,
)
from parser.heading_catalog import get_heading_rows


RELATIVE_HEADING_PATTERN = re.compile(r"^\d+\.?\s+.+")


def get_unique_sheet_name(name, used_names):
    base_name = safe_sheet_name(name)
    sheet_name = base_name
    counter = 1

    while sheet_name in used_names:
        suffix = f"_{counter}"
        sheet_name = f"{base_name[:31 - len(suffix)]}{suffix}"
        counter += 1

    used_names.add(sheet_name)
    return sheet_name


def format_issue_heading(issue):
    heading = issue["heading"]
    section_key = issue.get("section_key")
    if not section_key or not RELATIVE_HEADING_PATTERN.match(str(heading)):
        return heading

    parts = [part.strip() for part in str(section_key).split(">") if part.strip()]
    if len(parts) < 2:
        return heading

    return " > ".join(parts[-2:])


def write_issue_sheet(sheet, issues):
    """写入某个文档的问题明细。"""
    headers = ["序号", "问题类型", "所在章节", "位置", "内容", "问题说明"]
    sheet.append(headers)

    for col in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    issue_order = {
        "中文残留": 1,
        "厂商残留": 2,
        "术语问题": 3,
        "英文格式": 4,
        "版式问题": 5,
        "引用问题": 6,
        "内容质量": 7,
    }

    sorted_issues = sorted(
        issues,
        key=lambda x: issue_order.get(x["issue_type"], 999)
    )

    for index, issue in enumerate(sorted_issues, start=1):
        sheet.append([
            index,
            issue["issue_type"],
            format_issue_heading(issue),
            issue["location"],
            issue["content"],
            issue["issue"],
        ])

    widths = [8, 14, 30, 24, 90, 30]
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def apply_heading_catalog_style(sheet, headers, widths):
    """统一设置标题目录 Sheet 的表头、列宽和筛选。"""
    for col in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_heading_catalog_excel(all_file_items, output_file):
    """
    导出标题目录.xlsx。
    Sheet1：汇总
    后续 Sheet：每个文档一个标题目录 Sheet。
    """
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"

    summary_headers = ["序号", "文档名称", "标题数量"]
    summary_sheet.append(summary_headers)

    for col in range(1, len(summary_headers) + 1):
        cell = summary_sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    detail_headers = ["序号", "标题编号", "标题名称", "Heading级别", "位置", "完整标题"]
    detail_widths = [8, 16, 50, 16, 24, 70]
    used_sheet_names = {"汇总"}

    for file_index, (file_name, items) in enumerate(all_file_items.items(), start=1):
        heading_rows = get_heading_rows(items)
        summary_sheet.append([file_index, file_name, len(heading_rows)])

        sheet_name = get_unique_sheet_name(Path(file_name).stem, used_sheet_names)
        sheet = workbook.create_sheet(title=sheet_name)
        sheet.append(detail_headers)

        if heading_rows:
            for row_index, row in enumerate(heading_rows, start=1):
                sheet.append([
                    row_index,
                    row["heading_number"],
                    row["heading_title"],
                    row["heading_level"],
                    row["location"],
                    row["full_heading"],
                ])
        else:
            sheet.append([1, "", "未提取到标题", "", "", ""])

        apply_heading_catalog_style(sheet, detail_headers, detail_widths)

    summary_widths = [8, 50, 12]
    for i, width in enumerate(summary_widths, start=1):
        summary_sheet.column_dimensions[get_column_letter(i)].width = width

    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions

    workbook.save(output_file)


def export_summary_excel(all_file_issues, output_file):
    """
    导出总检查结果.xlsx：
    Sheet1：汇总
    后续Sheet：每个文档一个Sheet。
    """
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"

    summary_headers = ["序号", "文档名称", "问题总数", "中文残留", "厂商残留", "英文格式", "版式问题", "引用问题", "内容质量"]
    summary_sheet.append(summary_headers)

    for col in range(1, len(summary_headers) + 1):
        cell = summary_sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, (file_name, issues) in enumerate(all_file_issues.items(), start=1):
        chinese_count = sum(1 for i in issues if i["issue_type"] == "中文残留")
        vendor_count = sum(1 for i in issues if i["issue_type"] == "厂商残留")
        format_count = sum(1 for i in issues if i["issue_type"] == "英文格式")
        layout_count = sum(1 for i in issues if i["issue_type"] == "版式问题")
        reference_count = sum(1 for i in issues if i["issue_type"] == "引用问题")
        content_quality_count = sum(1 for i in issues if i["issue_type"] == "内容质量")
        summary_sheet.append([
            index,
            file_name,
            len(issues),
            chinese_count,
            vendor_count,
            format_count,
            layout_count,
            reference_count,
            content_quality_count,
        ])

    summary_widths = [8, 50, 12, 12, 12, 12, 12, 12, 12]
    for i, width in enumerate(summary_widths, start=1):
        summary_sheet.column_dimensions[get_column_letter(i)].width = width

    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions
    used_sheet_names = {"汇总"}

    for file_name, issues in all_file_issues.items():
        sheet_name = get_unique_sheet_name(Path(file_name).stem, used_sheet_names)
        sheet = workbook.create_sheet(title=sheet_name)

        if issues:
            write_issue_sheet(sheet, issues)
        else:
            headers = ["结果"]
            sheet.append(headers)
            sheet.cell(row=1, column=1).font = Font(bold=True)
            sheet.cell(row=1, column=1).fill = PatternFill("solid", fgColor="D9EAF7")
            sheet.append(["未发现中文残留、厂商残留或英文格式问题。"])
            sheet.column_dimensions["A"].width = 60

    workbook.save(output_file)


def _unpack_heading_check_result(sheet_name, result):
    if isinstance(result, dict):
        return result.get("document_name", sheet_name), result.get("issues", [])
    return sheet_name, result


def export_heading_check_report(all_results, output_dir):
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"

    summary_sheet.append(["文档", "问题数"])

    for sheet_name, result in all_results.items():
        document_name, issues = _unpack_heading_check_result(sheet_name, result)
        summary_sheet.append([document_name, len(issues)])

    used_sheet_names = {"汇总"}

    for sheet_name, result in all_results.items():
        _, issues = _unpack_heading_check_result(sheet_name, result)
        detail_sheet = workbook.create_sheet(
            get_unique_sheet_name(sheet_name, used_sheet_names)
        )
        detail_sheet.append(["类型", "编号", "原始内容", "完整标题", "说明"])

        if not issues:
            detail_sheet.append(["通过", "", "", "", "无问题"])
        else:
            for issue in issues:
                detail_sheet.append(issue)

    output_file = output_dir / get_heading_check_report_name()
    workbook.save(output_file)
    print("输出：", output_file)
    return output_file
