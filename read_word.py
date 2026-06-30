# 英语文档专用
print("我正在运行 read_word_heading_catalog_only_fixed.py")

from pathlib import Path
import re
from datetime import datetime
from collections import defaultdict

from docx import Document
from docx.oxml.ns import qn
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# 当前项目目录
BASE_DIR = Path(__file__).parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

SUMMARY_EXCEL_PREFIX = "总检查结果"
HEADING_CATALOG_EXCEL_PREFIX = "标题目录"


# ========== 通用工具 ==========

def clean_text(text):
    """清理文本中的多余换行和空格。"""
    if text is None:
        return ""
    text = text.replace("\u3000", " ")
    text = " ".join(text.split())
    return text.strip()


def get_summary_excel_name():
    """生成带时间戳的汇总 Excel 文件名，避免覆盖历史结果。"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{SUMMARY_EXCEL_PREFIX}_{timestamp}.xlsx"


def get_heading_catalog_excel_name():
    """生成带时间戳的标题目录 Excel 文件名，避免覆盖历史结果。"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{HEADING_CATALOG_EXCEL_PREFIX}_{timestamp}.xlsx"


def safe_sheet_name(name):
    """Excel Sheet 名不能超过31个字符，也不能包含部分特殊字符。"""
    invalid_chars = ["\\", "/", "?", "*", "[", "]", ":"]
    for ch in invalid_chars:
        name = name.replace(ch, "_")
    return name[:31] if len(name) > 31 else name


def is_heading(paragraph):
    """判断一个段落是否像章节标题。"""
    text = clean_text(paragraph.text)
    if not text:
        return False

    style_name = paragraph.style.name if paragraph.style else ""
    style_name_lower = style_name.lower()

    # Word 标题样式，例如 Heading 1 / Heading 2 / 标题 1
    if style_name_lower.startswith("heading") or style_name.startswith("标题"):
        return True

    # 手动输入的编号标题，例如 1 Overview / 1.1 Dashboard / 1.1.1 Technical Support
    if re.match(r"^\d+(\.\d+)*\s+.+", text):
        return True

    return False


def get_heading_level(paragraph):
    """返回 Word 标题样式层级：Heading 1 -> 1。"""
    style_name = paragraph.style.name if paragraph.style else ""
    match = re.match(r"^(?:Heading|标题)\s*(\d+)$", style_name, re.IGNORECASE)
    if not match:
        return None
    return int(match.group(1))

def is_toc_paragraph(paragraph):
    """识别目录段落；目录不参与正文标题层级检查。"""
    raw_text = paragraph.text or ""
    text = clean_text(raw_text)
    if not text:
        return False

    style_name = paragraph.style.name if paragraph.style else ""
    style_name_lower = style_name.lower()
    if style_name_lower.startswith("toc") or style_name.startswith("目录"):
        return True

    # Word 目录通常用制表符或点引导符连接页码。
    if re.match(r"^\d+(?:\.\d+)*\s+.+\t\d+$", raw_text):
        return True
    if re.match(r"^\d+(?:\.\d+)*\s+.+\.{2,}\s*\d+$", text):
        return True
    if re.match(r"^\d+(?:\.\d+)*\s+.+\s+\d+$", text):
        return True

    return False

# ========== Word 自动编号读取工具 ==========

def xml_val(element):
    """读取 XML 元素里的 w:val。"""
    if element is None:
        return None
    return element.get(qn("w:val"))


def int_or_default(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def get_paragraph_num_pr(paragraph):
    """
    获取段落的编号属性 numPr。
    先读段落自身的编号；如果没有，再读样式里的编号。
    这样可以识别 Word 标题样式绑定的自动编号。
    """
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.numPr is not None:
        return p_pr.numPr

    style = paragraph.style
    while style is not None:
        style_p_pr = style.element.pPr
        if style_p_pr is not None and style_p_pr.numPr is not None:
            return style_p_pr.numPr
        style = style.base_style

    return None


def get_paragraph_num_id_and_level(paragraph):
    """返回段落编号的 numId 和 ilvl。"""
    num_pr = get_paragraph_num_pr(paragraph)
    if num_pr is None:
        return None, None

    num_id = xml_val(num_pr.numId)
    ilvl = xml_val(num_pr.ilvl)

    # 有些样式只写 numId，不写 ilvl，标题默认按 0 级处理
    if num_id is not None and ilvl is None:
        ilvl = "0"

    return num_id, ilvl


def build_numbering_info(document):
    """
    解析 word/numbering.xml。
    返回：
    - num_to_abstract: numId -> abstractNumId
    - abstract_levels: abstractNumId -> ilvl -> 编号格式信息
    - num_overrides: numId -> ilvl -> 起始编号覆盖
    """
    num_to_abstract = {}
    abstract_levels = {}
    num_overrides = {}

    try:
        numbering = document.part.numbering_part.element
    except Exception:
        return num_to_abstract, abstract_levels, num_overrides

    # abstractNum 定义每一级编号长什么样，例如 %1.%2.%3
    for abstract_num in numbering.findall(qn("w:abstractNum")):
        abstract_id = abstract_num.get(qn("w:abstractNumId"))
        if abstract_id is None:
            continue

        abstract_levels[abstract_id] = {}
        for lvl in abstract_num.findall(qn("w:lvl")):
            ilvl = lvl.get(qn("w:ilvl"))
            if ilvl is None:
                continue

            start = 1
            num_fmt = "decimal"
            lvl_text = None

            start_el = lvl.find(qn("w:start"))
            if start_el is not None and xml_val(start_el) is not None:
                start = int_or_default(xml_val(start_el), 1)

            num_fmt_el = lvl.find(qn("w:numFmt"))
            if num_fmt_el is not None and xml_val(num_fmt_el):
                num_fmt = xml_val(num_fmt_el)

            lvl_text_el = lvl.find(qn("w:lvlText"))
            if lvl_text_el is not None and xml_val(lvl_text_el):
                lvl_text = xml_val(lvl_text_el)

            abstract_levels[abstract_id][ilvl] = {
                "start": start,
                "num_fmt": num_fmt,
                "lvl_text": lvl_text,
            }

    # num 把具体 numId 绑定到 abstractNumId
    for num in numbering.findall(qn("w:num")):
        num_id = num.get(qn("w:numId"))
        abstract_el = num.find(qn("w:abstractNumId"))
        abstract_id = xml_val(abstract_el)
        if num_id is not None and abstract_id is not None:
            num_to_abstract[num_id] = abstract_id

        for lvl_override in num.findall(qn("w:lvlOverride")):
            ilvl = lvl_override.get(qn("w:ilvl"))
            start_override_el = lvl_override.find(qn("w:startOverride"))
            if num_id is not None and ilvl is not None and start_override_el is not None:
                num_overrides.setdefault(num_id, {})[ilvl] = int_or_default(xml_val(start_override_el), 1)

    return num_to_abstract, abstract_levels, num_overrides


def number_to_letters(number, upper=True):
    """1 -> A/a, 2 -> B/b, 27 -> AA/aa。"""
    if number <= 0:
        return str(number)
    letters = []
    while number:
        number -= 1
        letters.append(chr(ord("A") + number % 26))
        number //= 26
    result = "".join(reversed(letters))
    return result if upper else result.lower()


def format_number(number, num_fmt):
    """按 Word 编号格式简单转换。常见标题一般都是 decimal。"""
    if num_fmt in ("decimal", "decimalZero"):
        return str(number).zfill(2) if num_fmt == "decimalZero" else str(number)
    if num_fmt == "upperLetter":
        return number_to_letters(number, upper=True)
    if num_fmt == "lowerLetter":
        return number_to_letters(number, upper=False)
    # roman 等格式不常用于你的手册标题，这里保守返回数字
    return str(number)


def make_number_string(num_id, ilvl, numbering_state):
    """
    根据 numId + ilvl 生成当前段落的自动编号字符串。
    numbering_state 用于记录每个编号列表当前计数。
    """
    if num_id is None or ilvl is None:
        return ""

    num_to_abstract = numbering_state["num_to_abstract"]
    abstract_levels = numbering_state["abstract_levels"]
    num_overrides = numbering_state["num_overrides"]
    counters = numbering_state["counters"]

    abstract_id = num_to_abstract.get(str(num_id))
    if abstract_id is None:
        return ""

    level = int_or_default(ilvl, 0)
    levels = abstract_levels.get(abstract_id, {})
    level_info = levels.get(str(level))
    if not level_info:
        return ""

    num_fmt = level_info.get("num_fmt", "decimal")
    lvl_text = level_info.get("lvl_text") or f"%{level + 1}."

    # bullet 不是章节编号，不返回
    if num_fmt == "bullet":
        return ""

    counter_key = str(num_id)
    if counter_key not in counters:
        counters[counter_key] = {}

    # 初始化计数。
    # 注意：如果当前先遇到的是 2 级/3 级标题，Word 实际显示时上级编号通常已经有值。
    # 旧逻辑会把所有上级都设为 start-1，导致 %1 被替换成 0，出现 0.1.2 这种错位。
    # 修复：上级默认设为 start；只有当前级别先设为 start-1 后再 +1。
    for i in range(level + 1):
        if i not in counters[counter_key]:
            start = levels.get(str(i), {}).get("start", 1)
            if str(i) in num_overrides.get(str(num_id), {}):
                start = num_overrides[str(num_id)][str(i)]
            counters[counter_key][i] = start - 1 if i == level else start

    counters[counter_key][level] += 1

    # 下级归零，避免 2.3 后面切到 2.4 时残留 2.3.5
    for i in list(counters[counter_key].keys()):
        if i > level:
            counters[counter_key].pop(i, None)

    number_string = lvl_text
    for i in range(9):
        placeholder = f"%{i + 1}"
        if placeholder not in number_string:
            continue
        value = counters[counter_key].get(i)
        level_fmt = levels.get(str(i), {}).get("num_fmt", "decimal")
        if value is None:
            value = levels.get(str(i), {}).get("start", 1)
        replacement = format_number(value, level_fmt)
        number_string = number_string.replace(placeholder, replacement)

    # 清理常见尾部符号，最后统一加空格
    return number_string.strip().rstrip(".").strip()


def get_heading_text(paragraph, numbering_state=None):
    """获取标题文本：优先拼接 Word 自动编号；失败则保留原标题文字。"""
    text = clean_text(paragraph.text)
    if not text:
        return "未识别到章节"

    if numbering_state is not None:
        num_id, ilvl = get_paragraph_num_id_and_level(paragraph)
        number = make_number_string(num_id, ilvl, numbering_state)
        if number:
            # 避免标题本身已经包含编号时重复，例如 2.9 Statistics -> 2.9 2.9 Statistics
            if not re.match(r"^" + re.escape(number) + r"(\s+|\.)", text):
                return f"{number} {text}"

    return text if text else "未识别到章节"


def iter_block_items(parent):
    """按 Word 原始顺序遍历正文中的段落和表格。"""
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc

    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


# ========== Word读取 ==========

def read_word_file(file_path):
    """
    读取 Word 文档内容。
    返回列表，每条结构：
    {
        file_name, heading, location, source_type, text
    }
    """
    document = Document(file_path)
    file_name = file_path.name
    results = []
    current_heading = "未识别到章节"
    paragraph_index = 0
    table_index = 0

    num_to_abstract, abstract_levels, num_overrides = build_numbering_info(document)
    numbering_state = {
        "num_to_abstract": num_to_abstract,
        "abstract_levels": abstract_levels,
        "num_overrides": num_overrides,
        "counters": {},
    }

    # 1. 按顺序读取正文段落和表格
    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            text = clean_text(block.text)
            if not text:
                continue

            paragraph_index += 1

            is_toc = is_toc_paragraph(block)
            heading_level = get_heading_level(block)
            if is_heading(block):
                current_heading = get_heading_text(block, numbering_state)
                results.append({
                    "file_name": file_name,
                    "heading": current_heading,
                    "location": f"正文-段落{paragraph_index}",
                    "source_type": "正文",
                    "text": text,
                    "is_heading": not is_toc,
                    "is_toc": is_toc,
                    "heading_level": heading_level,
                })
            else:
                results.append({
                    "file_name": file_name,
                    "heading": current_heading,
                    "location": f"正文-段落{paragraph_index}",
                    "source_type": "正文",
                    "text": text,
                    "is_heading": False,
                    "is_toc": False,
                    "heading_level": None,
                })

        elif isinstance(block, Table):
            table_index += 1
            for row_index, row in enumerate(block.rows, start=1):
                row_texts = []
                for cell in row.cells:
                    cell_text = clean_text(cell.text)
                    if cell_text:
                        row_texts.append(cell_text)

                if row_texts:
                    results.append({
                        "file_name": file_name,
                        "heading": current_heading,
                        "location": f"正文-表格{table_index}-第{row_index}行",
                        "source_type": "正文表格",
                        "text": " | ".join(row_texts),
                    })

    # 2. 读取页眉页脚：页眉页脚通常无法准确归属正文章节
    for section_index, section in enumerate(document.sections, start=1):
        for p_index, para in enumerate(section.header.paragraphs, start=1):
            text = clean_text(para.text)
            if text:
                results.append({
                    "file_name": file_name,
                    "heading": "页眉",
                    "location": f"第{section_index}节页眉-段落{p_index}",
                    "source_type": "页眉",
                    "text": text,
                })

        for t_index, table in enumerate(section.header.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                row_texts = []
                for cell in row.cells:
                    cell_text = clean_text(cell.text)
                    if cell_text:
                        row_texts.append(cell_text)
                if row_texts:
                    results.append({
                        "file_name": file_name,
                        "heading": "页眉",
                        "location": f"第{section_index}节页眉-表格{t_index}-第{row_index}行",
                        "source_type": "页眉表格",
                        "text": " | ".join(row_texts),
                    })

        for p_index, para in enumerate(section.footer.paragraphs, start=1):
            text = clean_text(para.text)
            if text:
                results.append({
                    "file_name": file_name,
                    "heading": "页脚",
                    "location": f"第{section_index}节页脚-段落{p_index}",
                    "source_type": "页脚",
                    "text": text,
                })

        for t_index, table in enumerate(section.footer.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                row_texts = []
                for cell in row.cells:
                    cell_text = clean_text(cell.text)
                    if cell_text:
                        row_texts.append(cell_text)
                if row_texts:
                    results.append({
                        "file_name": file_name,
                        "heading": "页脚",
                        "location": f"第{section_index}节页脚-表格{t_index}-第{row_index}行",
                        "source_type": "页脚表格",
                        "text": " | ".join(row_texts),
                    })

    return results


# ========== 检查规则 ==========

def has_chinese(text):
    """检查真实内容是否包含中文汉字。"""
    return re.search(r"[\u4e00-\u9fff]", text) is not None


FORBIDDEN_VENDOR_KEYWORDS = [
    "KSManage Ultra",
    "Inspur",
    "IEI",
    "IEISYSTEM",
    "MAGINFRA",
    "AIVRES",
    "浪潮",
    "至索",
    "inspur.com",
    "ieisystem.com",
    "magiops",
    "maginfra.com",
    "aivres.com",
    "@inspur.com",
    "@ieisystem.com",
    "@maginfra.com",
    "@aivres.com",
    "400-860-0011",
    "800-860-0011",
    "山东省济南市",
]


def check_vendor(text):
    """检查 KAYTUS 文档中的其他厂商残留。"""
    issues = []
    lower_text = text.lower()

    for keyword in FORBIDDEN_VENDOR_KEYWORDS:
        if keyword.lower() in lower_text:
            issues.append(f"发现厂商残留：{keyword}")

    return issues


def check_english_format(text):
    """检查常见英文格式问题。返回问题说明列表。"""
    issues = []

    # 真正的中文标点。注意：英文智能引号 “ ” ‘ ’ 暂时不作为错误。
    chinese_punctuation = [
        "，", "。", "：", "；",
        "（", "）",
        "【", "】",
        "《", "》",
    ]

    for p in chinese_punctuation:
        if p in text:
            issues.append(f"发现中文标点：{p}")

    # 连续两个以上空格
    if re.search(r" {2,}", text):
        issues.append("发现连续两个或多个空格")

    # 英文标点前多空格，例如 passing , / name : / save .
    if re.search(r"\s+[,.;:!?]", text):
        issues.append("英文标点前可能多了空格")

    # 英文标点后缺少空格，例如 failed.The / Switch:If
    # 避免误报小数、版本号、URL过多，这里只做简单检查。
    if re.search(r"[A-Za-z0-9][.!?:;][A-Z]", text):
        issues.append("英文标点后可能缺少空格")

    return issues


HEADING_NUMBER_PATTERN = re.compile(r"^(\d+(?:\.\d+)*)\b")
RELATIVE_LEVEL4_PATTERN = re.compile(r"^(\d+)\.?\s+.+")
MAX_HEADING_DEPTH = 4


def parse_heading_number(text):
    """提取标题开头的数字编号，例如 2.3.1。"""
    match = HEADING_NUMBER_PATTERN.match(clean_text(text))
    if not match:
        return None
    try:
        return tuple(int(part) for part in match.group(1).split("."))
    except ValueError:
        return None


def parse_relative_level4_number(text):
    """识别三级标题下形如“1. 标题”的相对四级编号。"""
    match = RELATIVE_LEVEL4_PATTERN.match(clean_text(text))
    if not match:
        return None
    return int(match.group(1))

def is_toc_like_heading(text):
    """跳过目录中常见的“编号 标题 页码”行，避免和正文标题混淆。"""
    text = clean_text(text)
    return re.match(r"^\d+(?:\.\d+)*\s+.+\s+\d+$", text) is not None


def make_heading_issue(item, issue):
    return {
        "file_name": item["file_name"],
        "issue_type": "标题层级",
        "heading": item["heading"],
        "location": item["location"],
        "content": item["heading"],
        "issue": issue,
    }


def format_heading_number(number):
    return ".".join(str(part) for part in number)


def number_from_style_level(level, style_counters):
    """按 Word Heading 样式层级重新生成标题编号。"""
    if level is None or level < 1 or level > MAX_HEADING_DEPTH:
        return None

    for upper_level in range(1, level):
        if style_counters[upper_level] == 0:
            style_counters[upper_level] = 1

    style_counters[level] += 1
    for lower_level in range(level + 1, MAX_HEADING_DEPTH + 1):
        style_counters[lower_level] = 0

    return tuple(style_counters[1:level + 1])


def check_heading_structure(items):
    """检查标题连续性、层级深度、重复编号和编号倒退。"""
    issues = []
    seen_numbers = {}
    max_number_by_parent = {}
    style_counters = [0] * (MAX_HEADING_DEPTH + 1)
    current_level3_number = None

    for item in items:
        if item.get("is_toc") or not item.get("is_heading"):
            continue

        heading_text = item["heading"]
        source_text = item.get("text", heading_text)
        if is_toc_like_heading(heading_text):
            continue

        heading_level = item.get("heading_level")
        heading_number = parse_heading_number(heading_text)
        explicit_number = parse_heading_number(source_text)
        relative_level4 = parse_relative_level4_number(heading_text)

        if heading_level is not None and heading_number is None:
            continue

        is_relative_level4 = False
        if heading_level == 4 and relative_level4 is not None and current_level3_number is not None:
            number = current_level3_number + (relative_level4,)
            is_relative_level4 = True
        elif heading_level is not None:
            number = number_from_style_level(heading_level, style_counters)
        elif relative_level4 is not None and current_level3_number is not None:
            number = current_level3_number + (relative_level4,)
        else:
            number = explicit_number

        if number is None:
            continue

        if len(number) == 3:
            current_level3_number = number
        elif len(number) < 3:
            current_level3_number = None

        number_text = format_heading_number(number)
        parent = number[:-1]
        current = number[-1]

        if len(number) > MAX_HEADING_DEPTH:
            issues.append(make_heading_issue(
                item,
                f"标题层级过深：{number_text}，当前最多建议 {MAX_HEADING_DEPTH} 级标题",
            ))

        if len(number) > 1 and parent not in seen_numbers and not is_relative_level4 and heading_level is None:
            parent_text = format_heading_number(parent)
            issues.append(make_heading_issue(
                item,
                f"标题层级跳跃：{number_text} 缺少上级标题 {parent_text}",
            ))

        if number in seen_numbers:
            issues.append(make_heading_issue(
                item,
                f"标题重复：{number_text} 已在 {seen_numbers[number]} 出现过",
            ))
            continue

        previous_max = max_number_by_parent.get(parent)
        if previous_max is None:
            if current != 1:
                expected = format_heading_number(parent + (1,))
                issues.append(make_heading_issue(
                    item,
                    f"标题连续性异常：当前为 {number_text}，该层级应从 {expected} 开始",
                ))
        elif current < previous_max:
            issues.append(make_heading_issue(
                item,
                f"标题倒退：同级标题已出现到 {format_heading_number(parent + (previous_max,))}，当前又出现 {number_text}",
            ))
        elif current > previous_max + 1:
            expected = format_heading_number(parent + (previous_max + 1,))
            issues.append(make_heading_issue(
                item,
                f"标题连续性异常：{format_heading_number(parent + (previous_max,))} 后出现 {number_text}，疑似缺少 {expected}",
            ))

        seen_numbers[number] = item["location"]
        max_number_by_parent[parent] = max(previous_max or 0, current)

    return issues

def check_items(items):
    """对一个文档读取结果进行检查。

    注意：这里不再执行标题连续性/标题层级判断，避免误报。
    标题只在“标题目录”中单独提取展示。
    """
    issues = []

    for item in items:
        text = item["text"]

        if has_chinese(text):
            issues.append({
                "file_name": item["file_name"],
                "issue_type": "中文残留",
                "heading": item["heading"],
                "location": item["location"],
                "content": text,
                "issue": "发现中文汉字",
            })

        vendor_issues = check_vendor(text)
        for issue in vendor_issues:
            issues.append({
                "file_name": item["file_name"],
                "issue_type": "厂商残留",
                "heading": item["heading"],
                "location": item["location"],
                "content": text,
                "issue": issue,
            })

        format_issues = check_english_format(text)
        for issue in format_issues:
            issues.append({
                "file_name": item["file_name"],
                "issue_type": "英文格式",
                "heading": item["heading"],
                "location": item["location"],
                "content": text,
                "issue": issue,
            })

    return issues


# ========== 导出结果 ==========

def export_read_result_to_txt(items, output_file):
    """每个文档单独输出全文读取结果，方便调试。"""
    with open(output_file, "w", encoding="utf-8") as f:
        for index, item in enumerate(items, start=1):
            f.write(f"{index}. 文档：{item['file_name']}\n")
            f.write(f"   章节：{item['heading']}\n")
            f.write(f"   位置：{item['location']}\n")
            f.write(f"   内容：{item['text']}\n")
            f.write("\n")


def write_issue_sheet(sheet, issues):
    """写入某个文档的问题明细。"""
    headers = ["序号", "问题类型", "所在章节", "位置", "内容", "问题说明"]
    sheet.append(headers)

    for col in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    issue_order = {
        "中文残留": 1,
        "厂商残留": 2,
        "术语问题": 3,
        "英文格式": 4,
    }

    sorted_issues = sorted(
        issues,
        key=lambda x: issue_order.get(x["issue_type"], 999)
    )

    for index, issue in enumerate(sorted_issues, start=1):
        sheet.append([
            index,
            issue["issue_type"],
            issue["heading"],
            issue["location"],
            issue["content"],
            issue["issue"],
        ])

    widths = [8, 14, 30, 24, 90, 30]
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions




def split_heading_number_and_title(heading_text):
    """把“1.2 Dashboard”拆成编号和标题名称；如果没有编号，则编号为空。"""
    text = clean_text(heading_text)
    match = re.match(r"^(\d+(?:\.\d+)*)\s+(.+)$", text)
    if match:
        return match.group(1), match.group(2)
    return "", text


def apply_heading_catalog_style(sheet, headers, widths):
    """统一设置标题目录 Sheet 的表头、列宽和筛选。"""
    for col in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def get_heading_rows(items):
    """从某个文档的读取结果中提取标题行。"""
    rows = []
    for item in items:
        if item.get("is_toc") or not item.get("is_heading"):
            continue

        heading_text = item.get("heading", "")
        heading_number, heading_title = split_heading_number_and_title(heading_text)
        heading_level = item.get("heading_level")
        heading_level_text = f"Heading {heading_level}" if heading_level else "未识别"

        rows.append({
            "heading_number": heading_number,
            "heading_title": heading_title,
            "heading_level": heading_level_text,
            "location": item.get("location", ""),
            "full_heading": heading_text,
        })

    return rows


def export_heading_catalog_excel(all_file_items, output_file):
    """
    导出标题目录.xlsx。
    Sheet1：汇总
    后续 Sheet：每个文档一个标题目录 Sheet。
    只提取标题清单，不判断标题是否连续，避免误报。
    """
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"

    summary_headers = ["序号", "文档名称", "标题数量"]
    summary_sheet.append(summary_headers)

    for col in range(1, len(summary_headers) + 1):
        cell = summary_sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    detail_headers = ["序号", "标题编号", "标题名称", "Heading级别", "位置", "完整标题"]
    detail_widths = [8, 16, 50, 16, 24, 70]

    for file_index, (file_name, items) in enumerate(all_file_items.items(), start=1):
        heading_rows = get_heading_rows(items)
        summary_sheet.append([file_index, file_name, len(heading_rows)])

        sheet_name = safe_sheet_name(Path(file_name).stem)
        sheet = workbook.create_sheet(title=sheet_name)
        sheet.append(detail_headers)

        if heading_rows:
            for row_index, row in enumerate(heading_rows, start=1):
                sheet.append([
                    row_index,
                    row["heading_number"],
                    row["heading_title"],
                    row["heading_level"],
                    row["location"],
                    row["full_heading"],
                ])
        else:
            sheet.append([1, "", "未提取到标题", "", "", ""])

        apply_heading_catalog_style(sheet, detail_headers, detail_widths)

    summary_widths = [8, 50, 12]
    for i, width in enumerate(summary_widths, start=1):
        summary_sheet.column_dimensions[get_column_letter(i)].width = width

    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions

    workbook.save(output_file)


def export_summary_excel(all_file_issues, output_file):
    """
    导出总检查结果.xlsx：
    Sheet1：汇总
    后续Sheet：每个文档一个Sheet
    注意：即使某个文档没有问题，也会创建对应Sheet。
    """
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"

    # 汇总Sheet
    summary_headers = ["序号", "文档名称", "问题总数", "中文残留", "厂商残留", "英文格式"]
    summary_sheet.append(summary_headers)

    for col in range(1, len(summary_headers) + 1):
        cell = summary_sheet.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, (file_name, issues) in enumerate(all_file_issues.items(), start=1):
        chinese_count = sum(1 for i in issues if i["issue_type"] == "中文残留")
        vendor_count = sum(1 for i in issues if i["issue_type"] == "厂商残留")
        format_count = sum(1 for i in issues if i["issue_type"] == "英文格式")
        summary_sheet.append([
            index,
            file_name,
            len(issues),
            chinese_count,
            vendor_count,
            format_count,
        ])

    summary_widths = [8, 50, 12, 12, 12, 12]
    for i, width in enumerate(summary_widths, start=1):
        summary_sheet.column_dimensions[get_column_letter(i)].width = width

    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions

    # 每个文档一个Sheet：即使没有问题，也创建Sheet
    for file_name, issues in all_file_issues.items():
        sheet_name = safe_sheet_name(Path(file_name).stem)
        sheet = workbook.create_sheet(title=sheet_name)

        if issues:
            write_issue_sheet(sheet, issues)
        else:
            headers = ["结果"]
            sheet.append(headers)
            sheet.cell(row=1, column=1).font = Font(bold=True)
            sheet.cell(row=1, column=1).fill = PatternFill("solid", fgColor="D9EAF7")
            sheet.append(["未发现中文残留、厂商残留或英文格式问题。"])
            sheet.column_dimensions["A"].width = 60

    workbook.save(output_file)

# ========== 主程序 ==========

def main():
    word_files = [
        file for file in INPUT_DIR.glob("*.docx")
        if not file.name.startswith("~$")
    ]

    if not word_files:
        print("input 文件夹里没有找到 .docx 文件")
        return

    # 保存每个文档的问题；即使0个问题，也保留这个文档
    all_file_issues = {}

    # 保存每个文档的标题目录；只提取，不自动判断连续性
    all_file_items = {}

    for word_file in word_files:
        print(f"正在读取：{word_file.name}")

        items = read_word_file(word_file)
        all_file_items[word_file.name] = items

        read_txt_file = OUTPUT_DIR / f"{word_file.stem}_读取结果.txt"
        export_read_result_to_txt(items, read_txt_file)
        print(f"读取结果已输出：{read_txt_file}")

        issues = check_items(items)
        all_file_issues[word_file.name] = issues

        print(f"检查完成：{word_file.name}，发现 {len(issues)} 个问题")

    summary_excel_file = OUTPUT_DIR / get_summary_excel_name()
    export_summary_excel(all_file_issues, summary_excel_file)
    print(f"总检查结果已输出：{summary_excel_file}")

    heading_catalog_file = OUTPUT_DIR / get_heading_catalog_excel_name()
    export_heading_catalog_excel(all_file_items, heading_catalog_file)
    print(f"标题目录已输出：{heading_catalog_file}")

if __name__ == "__main__":
    main()
























