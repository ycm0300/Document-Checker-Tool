from openpyxl import load_workbook, Workbook
from pathlib import Path
from datetime import datetime
import re


BASE_DIR = Path(__file__).parent
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

CATALOG_PREFIX = "标题目录"
REPORT_PREFIX = "标题检查结果"


def clean(v):
    if v is None:
        return ""
    return str(v).strip()


def parse_num(text):
    """
    提取纯编号：2.1 / 6.4.1
    """
    text = clean(text)
    m = re.match(r"^(\d+(?:\.\d+)*)", text)
    if not m:
        return None
    return tuple(map(int, m.group(1).split(".")))


def find_latest():
    files = list(OUTPUT_DIR.glob(f"{CATALOG_PREFIX}*.xlsx"))
    if not files:
        files = list(BASE_DIR.glob(f"{CATALOG_PREFIX}*.xlsx"))
    return max(files, key=lambda p: p.stat().st_mtime)


def check_sheet(ws):
    headers = [clean(c.value) for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}

    if "标题编号" not in idx:
        raise ValueError(f"{ws.title} 缺少：标题编号")

    col_num = idx["标题编号"]

    seen = set()
    issues = []

    last_level = {}

    rows = []

    # 收集数据
    for r in range(2, ws.max_row + 1):
        val = ws.cell(r, col_num).value
        num = parse_num(val)
        if not num:
            continue

        rows.append((r, num, val))

    # =========================
    # 纯编号分析
    # =========================
    for i, (r, num, raw) in enumerate(rows):

        # -------- 重复检查 --------
        if num in seen:
            issues.append([
                "重复",
                ".".join(map(str, num)),
                raw,
                f"第{r}行重复出现"
            ])
            continue

        seen.add(num)

        # -------- 连续性检查 --------
        parent = num[:-1]
        cur = num[-1]

        if parent in last_level:
            expect = last_level[parent] + 1
            if cur != expect:
                issues.append([
                    "不连续",
                    ".".join(map(str, num)),
                    raw,
                    f"应为 {'.'.join(map(str, parent + (expect,)))}"
                ])

        last_level[parent] = cur

    return issues


def write_report(all_results):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "汇总"

    ws0.append(["文档", "问题数"])

    for sheet, issues in all_results.items():
        ws0.append([sheet, len(issues)])

    for sheet, issues in all_results.items():
        ws = wb.create_sheet(sheet[:31])
        ws.append(["类型", "编号", "原始内容", "说明"])

        if not issues:
            ws.append(["通过", "", "", "无问题"])
        else:
            for i in issues:
                ws.append(i)

    out = OUTPUT_DIR / f"{REPORT_PREFIX}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(out)
    print("输出：", out)


def main():
    file = find_latest()
    wb = load_workbook(file)

    results = {}

    for sheet in wb.sheetnames:
        if sheet == "汇总":
            continue

        ws = wb[sheet]
        results[sheet] = check_sheet(ws)

    write_report(results)


if __name__ == "__main__":
    main()