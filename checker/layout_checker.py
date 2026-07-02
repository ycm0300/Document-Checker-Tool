from docx import Document
from docx.oxml.ns import qn

from config.common_rules import clean_text
from parser.word_reader import get_heading_text, get_heading_level, is_toc_paragraph


def _style_page_break_before(style):
    while style is not None:
        if style.paragraph_format.page_break_before:
            return True
        style = style.base_style
    return False


def _paragraph_has_page_break(paragraph):
    for br in paragraph._p.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _paragraph_has_next_page_section_break(paragraph):
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.sectPr is None:
        return False

    section_type = p_pr.sectPr.find(qn("w:type"))
    if section_type is None:
        return True

    return section_type.get(qn("w:val")) != "continuous"


def _paragraph_starts_new_page(paragraph):
    if paragraph.paragraph_format.page_break_before:
        return True
    if paragraph.style is not None and _style_page_break_before(paragraph.style):
        return True
    if _paragraph_has_page_break(paragraph):
        return True
    return False


def _has_content_before(paragraphs, heading_index):
    for paragraph in paragraphs[:heading_index]:
        if clean_text(paragraph.text):
            return True
    return False


def _find_previous_layout_boundary(paragraphs, heading_index):
    blank_count = 0

    for index in range(heading_index - 1, -1, -1):
        paragraph = paragraphs[index]
        if clean_text(paragraph.text):
            return paragraph, blank_count

        if _paragraph_has_page_break(paragraph) or _paragraph_has_next_page_section_break(paragraph):
            return paragraph, blank_count

        blank_count += 1

    return None, blank_count


def _make_issue(file_name, heading, paragraph_index, issue):
    return {
        "file_name": file_name,
        "issue_type": "版式问题",
        "heading": heading,
        "location": f"正文-段落{paragraph_index}",
        "content": heading,
        "issue": issue,
    }


def check_heading1_starts_new_page(file_path):
    """检查 Heading 1 / 标题 1 是否设置为新页开始。"""
    document = Document(file_path)
    issues = []

    for index, paragraph in enumerate(document.paragraphs):
        text = clean_text(paragraph.text)
        if not text or is_toc_paragraph(paragraph):
            continue

        if get_heading_level(paragraph) != 1:
            continue

        if not _has_content_before(document.paragraphs, index):
            continue

        if _paragraph_starts_new_page(paragraph):
            continue

        previous_boundary, blank_count = _find_previous_layout_boundary(document.paragraphs, index)
        has_boundary_page_break = (
            previous_boundary is not None
            and (
                _paragraph_has_page_break(previous_boundary)
                or _paragraph_has_next_page_section_break(previous_boundary)
            )
        )

        if has_boundary_page_break and blank_count > 0:
            issues.append(_make_issue(
                file_path.name,
                get_heading_text(paragraph),
                index + 1,
                "一级标题前存在分页符或分节符，但与标题之间有多余空段落",
            ))
            continue

        if has_boundary_page_break:
            continue

        issues.append(_make_issue(
            file_path.name,
            get_heading_text(paragraph),
            index + 1,
            "一级标题未检测到分页符、分节符或段落分页前设置",
        ))

    return issues
