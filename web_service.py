from pathlib import Path

from openpyxl import load_workbook

from checker.content_quality_checker import check_duplicate_paragraphs
from checker.layout_checker import check_heading1_starts_new_page
from checker.reference_checker import check_figure_references
from main import check_items
from parser.word_reader import read_word_file
from report.excel_writer import export_summary_excel


def check_word_file(word_file, progress_callback=None):
    """检查单个 Word 文件并返回问题列表。"""
    word_file = Path(word_file)
    if progress_callback:
        progress_callback(0, "正在读取文档")
    items = read_word_file(word_file)
    if progress_callback:
        progress_callback(0.2, "正在检查文本格式")
    issues = check_items(items)
    if progress_callback:
        progress_callback(0.6, "正在检查重复内容")
    issues.extend(check_duplicate_paragraphs(items))
    if progress_callback:
        progress_callback(0.72, "正在检查图片引用")
    issues.extend(check_figure_references(items))
    if progress_callback:
        progress_callback(0.84, "正在检查标题版式")
    issues.extend(check_heading1_starts_new_page(word_file))
    return issues


def check_word_files(word_files, output_dir, progress_callback=None):
    """检查网页上传的 Word 文件并生成一个多 Sheet Excel。"""
    all_file_issues = {}
    total = len(word_files)
    for index, word_file in enumerate(word_files):
        word_file = Path(word_file)
        def report_phase(fraction, message):
            if progress_callback:
                progress_callback(index + fraction, total, word_file.name, message)

        all_file_issues[word_file.name] = check_word_file(word_file, report_phase)

    if progress_callback:
        progress_callback(total, total, None, "文档检查完成，正在生成报告")
    excel_file = export_web_excel(all_file_issues, output_dir)
    return excel_file, all_file_issues


def export_web_excel(all_file_issues, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    excel_file = output_dir / "总检查结果.xlsx"
    export_summary_excel(all_file_issues, excel_file)
    return excel_file


def workbook_to_sheets(excel_file):
    """将 Excel 的所有 Sheet 转成适合前端展示的二维数组。"""
    workbook = load_workbook(excel_file, data_only=True, read_only=True)
    try:
        return [
            {
                "name": sheet.title,
                "rows": [
                    ["" if value is None else str(value) for value in row]
                    for row in sheet.iter_rows(values_only=True)
                ],
            }
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()
